"""Rewrite the SIH deck's content to match what the system actually does and measures.

Edits happen run-by-run so the template's fonts, sizes, colours and shape geometry survive
untouched: only the strings inside existing runs change. Replacement text is kept at or below
the length of what it replaces wherever possible, because this machine has no LibreOffice and
so the result cannot be rendered and checked — a longer line could overflow its box unseen.
"""
import copy

from pptx import Presentation
from pptx.util import Inches

SRC = "/Users/adi/sih-2026/SIH2026-IDEA-Presentation-Format.pptx"
OUT = "/Users/adi/sih-2026/SIH2026-IDEA-Presentation-Meridian.pptx"

prs = Presentation(SRC)
slides = list(prs.slides)
report: list[str] = []


def shape(slide_no: int, shape_id: int):
    for sh in slides[slide_no - 1].shapes:
        if sh.shape_id == shape_id:
            return sh
    raise KeyError(f"slide {slide_no} has no shape {shape_id}")


def set_run(slide_no: int, shape_id: int, para: int, run: int, new: str) -> None:
    r = shape(slide_no, shape_id).text_frame.paragraphs[para].runs[run]
    old = r.text
    r.text = new
    flag = "  ← LONGER" if len(new) > len(old) else ""
    report.append(f"  s{slide_no}/{shape_id} p{para}r{run}  {len(old):>3}→{len(new):<3}{flag}\n"
                  f"      - {old}\n      + {new}")


# ── Slide 2 · the benchmark result, which is the strongest line we have ─────────────────
#
# The empty last paragraph of the uniqueness box is reused rather than a new one appended, so
# the box's paragraph count and therefore its height are unchanged.
box = shape(2, 108).text_frame
donor = box.paragraphs[3]
target = box.paragraphs[4]
for r in list(target.runs):
    r._r.getparent().remove(r._r)
for src_run in donor.runs:                       # carry the bullet's exact formatting across
    target._p.append(copy.deepcopy(src_run._r))
set_run(2, 108, 4, 0, "•")
set_run(2, 108, 4, 1, "Benchmarked: no fuzzy or embedding baseline reaches our precision")

# The cascade feature can carry its own measurement without growing.
set_run(2, 107, 7, 1,
        " Identity to attribute rules to model, only on ambiguous pairs: 99.4% of pairs cut.")

# ── Slide 3 · the stack, which currently names four things we do not use ────────────────
set_run(3, 125, 0, 0,
        "React 19, Vite, TypeScript, TanStack Query, Tailwind CSS 4, Recharts, Axios")
set_run(3, 133, 0, 1,
        ", REST APIs, FastAPI, SQLAlchemy, SQLite (Postgres-ready), Pydantic, pytest (75)")
set_run(3, 136, 0, 0, "Hugging Face [BGE-base],")
set_run(3, 136, 0, 1, " Qwen2.5-7B (Ollama), NumPy, PyYAML, ")

# Drop the Pandas and scikit-learn marks, then re-centre what is left in the AI/ML panel so
# the row does not read as two gaps where something was deleted.
PANEL_LEFT, PANEL_WIDTH = Inches(8.90), Inches(4.21)
for dead in (144, 146):
    sh = shape(3, dead)
    sh._element.getparent().remove(sh._element)
    report.append(f"  s3/{dead}  logo removed")

kept = [shape(3, i) for i in (142, 143, 145)]
gap = Inches(0.35)
total = sum(s.width for s in kept) + gap * (len(kept) - 1)
x = PANEL_LEFT + (PANEL_WIDTH - total) // 2
for sh in kept:
    sh.left = int(x)
    x += sh.width + gap
report.append(f"  s3  re-centred {len(kept)} remaining AI/ML logos")

# ── Slide 4 · feasibility claims, each given the number that backs it ───────────────────
set_run(4, 159, 1, 1, " - All 7,117 pairs decided offline, zero model calls.")
set_run(4, 159, 2, 1, " - Blocking cuts 99.4% of comparisons, missing no true pair.")
set_run(4, 159, 4, 1, " - Dry run on a CPSE master: 0 existing fields altered.")
set_run(4, 159, 5, 1, "- Gates outside the AI: 0.0028 false-merge rate on hard pairs.")
set_run(4, 159, 6, 1, " - 4 families, each added as YAML with no code change.")

# Cost viability should describe our cost, not somebody else's loss. The CAG figure moves to
# slide 5, where it is already correctly framed as reference scale rather than as savings.
set_run(4, 160, 2, 0,
        "Cost viability - The deterministic path needs no model call, so most decisions "
        "cost nothing to run.")

# ── Slide 5 · drop the second causal use of the CAG number, upgrade the audit claim ─────
set_run(5, 176, 1, 1,
        " One identity per material lets CPSEs aggregate demand and compare prices.")
set_run(5, 176, 2, 1, " SHA-256 hash-chained audit trail; 1,935 events, tamper-evident.")

prs.save(OUT)
print("\n".join(report))
print(f"\nsaved {OUT}")
